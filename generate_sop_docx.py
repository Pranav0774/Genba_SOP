import os
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

VIDEO_NAME = "eng_05"

STEPS_PATH = rf"D:\genba-sop\reference\English\SOPs\{VIDEO_NAME}_steps.txt"
TOOLS_SAFETY_PATH = rf"D:\genba-sop\reference\English\SOPs\{VIDEO_NAME}_tools_safety.txt"
PHOTOS_FOLDER = rf"D:\genba-sop\photos\English\{VIDEO_NAME}"
SCENE_CSV_PATH = rf"D:\genba-sop\photos\{VIDEO_NAME}\{VIDEO_NAME}-Scenes.csv"
OUTPUT_PATH = rf"D:\genba-sop\reference\English\SOPs\{VIDEO_NAME}_SOP.docx"
LOG_PATH = rf"D:\genba-sop\reference\English\SOPs\{VIDEO_NAME}_selected_images_log.txt"

EXCLUDE_PHOTOS_BEFORE_SECONDS = {
    "eng_04": 64.5,
}.get(VIDEO_NAME, 0.0)


def parse_steps(path):
    with open(path, "r", encoding="utf-8") as f:
        lines = f.read().strip().split("\n")

    steps = []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        match = re.match(r"^\d+\.\s*\[(\d+(?:\.\d+)?)s?\]\s*(.+)$", line)
        if match:
            timestamp = float(match.group(1))
            text = match.group(2)
            is_unclear = "[UNCLEAR]" in text
            text = text.replace("[UNCLEAR]", "").strip()
            steps.append((text, is_unclear, timestamp))
        else:
            match2 = re.match(r"^\d+\.\s*(.+)$", line)
            if match2:
                text = match2.group(1)
                is_unclear = "[UNCLEAR]" in text
                text = text.replace("[UNCLEAR]", "").strip()
                steps.append((text, is_unclear, None))
    return steps


def parse_tools_safety(path):
    with open(path, "r", encoding="utf-8") as f:
        content = f.read()

    tools_match = re.search(r"TOOLS AND MATERIALS\s*(.*?)\s*SAFETY POINTS", content, re.DOTALL)
    safety_match = re.search(r"SAFETY POINTS\s*(.*)", content, re.DOTALL)

    tools = []
    if tools_match:
        tools = [l.strip("- ").strip() for l in tools_match.group(1).split("\n") if l.strip()]

    safety = []
    if safety_match:
        safety = [l.strip("- ").strip() for l in safety_match.group(1).split("\n") if l.strip()]

    return tools, safety


def get_candidate_photos(folder, max_photos=None):
    photos = sorted([
        os.path.join(folder, f) for f in os.listdir(folder)
        if f.lower().endswith((".jpg", ".jpeg", ".png"))
    ])
    if max_photos:
        step = max(1, len(photos) // max_photos)
        photos = photos[::step][:max_photos]
    return photos


def load_scene_timestamps(csv_path):
    import csv
    scene_times = {}
    with open(csv_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    reader = csv.DictReader(lines[1:])
    for row in reader:
        try:
            scene_num = int(row["Scene Number"])
            start_sec = float(row["Start Time (seconds)"])
            end_sec = float(row["End Time (seconds)"])
            scene_times[scene_num] = (start_sec, end_sec)
        except (KeyError, ValueError):
            continue
    return scene_times


def get_photo_scene_number(photo_filename):
    match = re.search(r"Scene-(\d+)", photo_filename)
    return int(match.group(1)) if match else None


def get_photo_position(photo_filename):
    match = re.search(r"Scene-\d+-(\d+)", photo_filename)
    return int(match.group(1)) if match else 1


def build_scene_position_counts(photos):
    counts = {}
    for p in photos:
        scene_num = get_photo_scene_number(os.path.basename(p))
        position = get_photo_position(os.path.basename(p))
        if scene_num is not None:
            counts[scene_num] = max(counts.get(scene_num, 1), position)
    return counts


def get_photo_timestamp(photo_filename, scene_times, scene_position_counts):
    scene_num = get_photo_scene_number(photo_filename)
    if scene_num is None or scene_num not in scene_times:
        return None
    start, end = scene_times[scene_num]
    position = get_photo_position(photo_filename)
    total = scene_position_counts.get(scene_num, 3)
    if total <= 1:
        return (start + end) / 2
    fraction = (position - 1) / (total - 1)
    fraction = max(0.0, min(1.0, fraction))
    return start + fraction * (end - start)


def match_photos_to_steps(steps, photos, scene_times, exclude_before_seconds=0.0):
    scene_position_counts = build_scene_position_counts(photos)
    photo_ts = []
    for p in photos:
        ts = get_photo_timestamp(os.path.basename(p), scene_times, scene_position_counts)
        if ts is not None and ts < exclude_before_seconds:
            continue
        photo_ts.append((p, ts))

    have_any_timestamps = any(ts is not None for _, ts in photo_ts)

    matched = [None] * len(steps)
    if have_any_timestamps:
        available = [(p, ts) for p, ts in photo_ts if ts is not None]

        pairs = []
        for step_idx, (step_text, is_unclear, step_ts) in enumerate(steps):
            if step_ts is None:
                continue
            for photo_path, photo_t in available:
                pairs.append((abs(photo_t - step_ts), step_idx, photo_path))
        pairs.sort(key=lambda x: x[0])

        assigned_steps = set()
        used_photos = set()
        for _, step_idx, photo_path in pairs:
            if step_idx in assigned_steps or photo_path in used_photos:
                continue
            matched[step_idx] = photo_path
            assigned_steps.add(step_idx)
            used_photos.add(photo_path)

        for step_idx, (step_text, is_unclear, step_ts) in enumerate(steps):
            if step_ts is None or matched[step_idx] is not None:
                continue
            if available:
                best = min(available, key=lambda x: abs(x[1] - step_ts))
                matched[step_idx] = best[0]
    else:
        num_steps = len(steps)
        num_photos = len(photos)
        for i in range(num_steps):
            if num_photos > 0:
                idx = min(int(i * num_photos / num_steps), num_photos - 1)
                matched[i] = photos[idx]

    return matched


def build_document(video_name, steps, tools, safety, matched_photos, output_path, log_path):
    doc = Document()

    title = doc.add_heading(f"Standard Operating Procedure — {video_name}", level=0)

    doc.add_heading("Tools and Materials", level=1)
    if tools:
        for tool in tools:
            doc.add_paragraph(tool, style="List Bullet")
    else:
        doc.add_paragraph("None listed.")

    doc.add_heading("Safety Points", level=1)
    if safety:
        for point in safety:
            doc.add_paragraph(point, style="List Bullet")
    else:
        doc.add_paragraph("None mentioned in transcript.")

    doc.add_heading("Procedure Steps", level=1)

    log_lines = ["Selected Image Log — {}".format(video_name), "=" * 60]

    for i, ((step_text, is_unclear, step_ts), photo_path) in enumerate(zip(steps, matched_photos), 1):
        p = doc.add_paragraph()
        ts_label = f" (at {step_ts:.1f}s)" if step_ts is not None else ""
        run = p.add_run(f"{i}. {step_text}{ts_label}")
        run.font.size = Pt(11)

        if is_unclear:
            flag = p.add_run("  要確認")
            flag.bold = True

        photo_name = os.path.basename(photo_path) if photo_path else "NO MATCH FOUND"
        if photo_path:
            try:
                doc.add_picture(photo_path, width=Inches(3.5))
                caption = doc.add_paragraph()
                caption_run = caption.add_run(f"Selected image: {photo_name}")
                caption_run.italic = True
                caption_run.font.size = Pt(9)
            except Exception as e:
                doc.add_paragraph(f"[Could not embed photo: {e}]")
        else:
            doc.add_paragraph("[No matching candidate photo found for this step's timestamp]")

        log_lines.append(
            f"Step {i} [{step_ts if step_ts is not None else 'N/A'}s]: {step_text}"
        )
        log_lines.append(f"  -> Selected image: {photo_name}")
        log_lines.append("")

    doc.save(output_path)
    print(f"Saved SOP document to: {output_path}")

    with open(log_path, "w", encoding="utf-8") as f:
        f.write("\n".join(log_lines))
    print(f"Saved selected-image log to: {log_path}")


if __name__ == "__main__":
    print("Parsing steps...")
    steps = parse_steps(STEPS_PATH)
    print(f"Found {len(steps)} steps ({sum(1 for _, u, _ in steps if u)} marked unclear)")

    print("Parsing tools and safety...")
    tools, safety = parse_tools_safety(TOOLS_SAFETY_PATH)
    print(f"Found {len(tools)} tools, {len(safety)} safety points")

    print("Loading candidate photos...")
    photos = get_candidate_photos(PHOTOS_FOLDER)
    print(f"Found {len(photos)} candidate photos")

    print("Loading scene timestamps...")
    scene_times = load_scene_timestamps(SCENE_CSV_PATH)
    print(f"Loaded {len(scene_times)} scene timestamps")

    print("Matching photos to steps by timestamp...")
    matched_photos = match_photos_to_steps(steps, photos, scene_times, EXCLUDE_PHOTOS_BEFORE_SECONDS)
    matched_count = sum(1 for p in matched_photos if p is not None)
    print(f"Matched {matched_count}/{len(steps)} steps to a photo")

    print("Building Word document...")
    build_document(VIDEO_NAME, steps, tools, safety, matched_photos, OUTPUT_PATH, LOG_PATH)
